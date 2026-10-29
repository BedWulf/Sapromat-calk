"""Отчёт по методичке ОмГТУ, задача №2 «Кручение» (2026/27)."""
from __future__ import annotations
from docx import Document
from docx.shared import Cm,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from model import Calculation
from teacher import strength_design,DATA
from teacher_plot import teacher_png_bytes
import io
import math

def fmt(x,n=3):
    if abs(x)<.5*10**(-n):x=0.
    return f'{x:.{n}f}'.rstrip('0').rstrip('.').replace('.',',')

def add(d,text='',bold=False,align=None,space=4):
    p=d.add_paragraph()
    if align is not None:p.alignment=align
    p.paragraph_format.line_spacing=1.15
    p.paragraph_format.space_after=Pt(space)
    r=p.add_run(text);r.bold=bold;r.font.name='Times New Roman';r.font.size=Pt(11.5)
    return p

def h(d,text):return add(d,text,True,space=5)

def table(d,rows,width=None,size=9.7):
    t=d.add_table(rows=0,cols=len(rows[0]))
    t.style='Table Grid'
    t.autofit=False
    for ri,items in enumerate(rows):
        cells=t.add_row().cells
        for j,item in enumerate(items):
            cells[j].text=str(item)
            if width:cells[j].width=Cm(width[j])
            cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[j].paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(size);r.bold=(ri==0)
    d.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def pg(d):d.add_page_break()

def create_report(c:Calculation,path:str):
    s=c.s
    des=strength_design(c)
    d=Document()
    sec=d.sections[0];sec.page_height=Cm(29.7);sec.page_width=Cm(21)
    sec.top_margin=Cm(2);sec.bottom_margin=Cm(2);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
    normal=d.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11.5)
    normal.paragraph_format.space_after=Pt(4)
    center=WD_ALIGN_PARAGRAPH.CENTER
    # Титульный лист повторяет текст и последовательность образца преподавателя.
    add(d,'Федеральное государственное автономное образовательное учреждение',align=center,space=0)
    add(d,'высшего образования «Омский государственный технический университет»',align=center,space=58)
    add(d,'Кафедра «Машиноведение»',align=center,space=80)
    add(d,'ДОМАШНЕЕ ЗАДАНИЕ № 2',True,center,22)
    add(d,'по дисциплине',align=center,space=7)
    add(d,'«СОПРОТИВЛЕНИЕ МАТЕРИАЛОВ»',True,center,15)
    add(d,'за I семестр 2026/27 уч. года',align=center,space=3)
    add(d,'по теме «Кручение»',align=center,space=26)
    add(d,'Вариант: '+(s.title_variant or '________________'),align=center,space=56)
    for t in [
        'Выполнил: ст. гр. '+(s.title_group or '____________'),
        s.title_student or 'Ф.И.О. _______________',
        'Дата: '+(s.title_date or '____________'),
        'Проверил: Ст. преп. каф.',
        '«Машиноведение»',
        (s.title_teacher or 'Онищук С.Ю.')+' __________',
        'Дата: __________'
    ]:
        p=add(d,t,space=4);p.paragraph_format.left_indent=Cm(7.1)
    add(d,'Омск – '+(s.title_year or '2026'),align=center,space=0)
    pg(d)
    h(d,'ЗАДАЧА 2. Кручение')
    add(d,'Дано: значения нагрузок, схема нагружения, способ закрепления и координаты приложения нагрузок. Для задачи принята длина участка l = 1 м.' if s.teacher_scheme else 'Дано: собственная схема нагружения, нагрузки и размеры участков, введённые пользователем.')
    if s.teacher_scheme:
        L,m=DATA[s.teacher_data_row]
        table(d,[['№ строки','№ схемы','L, кН·м','m, кН·м/м','l, м','[τ], МПа','d/D'],
               [s.teacher_data_row,s.teacher_scheme,fmt(L),fmt(m),'1',fmt(s.tau_allow_MPa),'0,5']],size=9)
    else:
        table(d,[['Закрепление','[τ], МПа','Длины участков, м','G, МПа'],
           ['слева' if s.fixed=='left' else 'справа',fmt(s.tau_allow_MPa),', '.join(map(fmt,s.lengths_m)),fmt(s.G_MPa)]],size=8.7)
        table(d,[['Участок','I','II','III','IV'],
          ['m, кН·м/м']+[fmt(x) for x in s.distributed_knm_per_m],
          ['l, м']+[fmt(x) for x in s.lengths_m]],size=9)
        table(d,[['Узел','0','1','2','3','4'],['L, кН·м']+[fmt(x) for x in s.node_moments_knm]],size=9)
    h(d,'Требуется:')
    add(d,'1. Записать уравнение крутящего момента Mк(z), граничные условия задачи, рассчитать восемь значений крутящего момента по участкам.')
    add(d,'2. Выполнить чертёж схемы и эпюру в масштабе. Определить опасное сечение.')
    add(d,'3. Из условия прочности подобрать размеры круглого и кольцевого поперечных сечений при D:d = 2:1, приняв [τ] = '+fmt(s.tau_allow_MPa)+' МПа. Сравнить веса выбранных стержней.')
    h(d,'Решение:')
    add(d,'1. Совместим начало координат «0» с левым концом вала и направим ось z вдоль вала. Разделим его на четыре участка I–IV.')
    add(d,'2. Внешние моменты проецируем на ось z по правилу буравчика: точка сверху соответствует положительному внешнему моменту, знак «+» сверху — отрицательному. Внутренний Mк(z) равен сумме внешних моментов слева от сечения, взятой с обратным знаком.')
    add(d,'3. Из условия равновесия ΣM_z = 0 получаем реактивный момент заделки: M_R = −ΣM_внеш = −('+fmt(c.applied_sum)+') = '+fmt(c.reaction_knm)+' кН·м.')
    boundary='Mк(4l) = '+fmt(c.end_T[-1]) if s.fixed=='left' else 'Mк(0) = '+fmt(c.start_T[0])
    derived='Mк(0) = '+fmt(c.start_T[0]) if s.fixed=='left' else 'Mк(4l) = '+fmt(c.end_T[-1])
    add(d,'4. Граничное условие на свободном конце: '+boundary+' кН·м. Противоположное условие определяется из равновесия: '+derived+' кН·м.')
    add(d,'Проверка: ΣM_внеш + M_R = '+fmt(c.applied_sum)+' + ('+fmt(c.reaction_knm)+') = 0 кН·м.')
    pg(d)
    h(d,'5. Уравнения крутящих моментов по участкам')
    add(d,'На каждом участке Mк(z) линейно меняется под действием распределённых моментов. В точках приложения сосредоточенных моментов на эпюре возникает скачок, поэтому значение на конце предыдущего участка и в начале следующего может различаться.')
    roman=['I','II','III','IV']
    for i in range(4):
        a,b=c.coords[i:i+2];T=c.start_T[i];q=s.distributed_knm_per_m[i]
        add(d,roman[i]+' участок: '+fmt(a)+' ≤ z ≤ '+fmt(b)+' м.',True,space=2)
        add(d,'Mк(z) = '+fmt(T)+' − ('+fmt(q)+')·(z − '+fmt(a)+'), кН·м.',space=2)
        add(d,'Mк('+fmt(a)+') = '+fmt(T)+' кН·м;  Mк('+fmt(b)+') = '+fmt(c.end_T[i])+' кН·м.',space=4)
    h(d,'6. Таблица восьми значений')
    table(d,[['Участок','I','I','II','II','III','III','IV','IV'],
         ['z, м']+[fmt(v) for i in range(4) for v in (c.coords[i],c.coords[i+1])],
         ['Mк, кН·м']+[fmt(v) for i in range(4) for v in (c.start_T[i],c.end_T[i])]],size=8.5)
    worst=[]
    for i in range(4):
        if abs(c.start_T[i])>=c.max_T_knm-1e-8:worst.append(fmt(c.coords[i]))
        if abs(c.end_T[i])>=c.max_T_knm-1e-8:worst.append(fmt(c.coords[i+1]))
    add(d,'Опасное сечение: z = '+', '.join(dict.fromkeys(worst))+' м; |Mк|max = '+fmt(c.max_T_knm)+' кН·м.')
    pg(d)
    h(d,'7. Схема нагружения и эпюра крутящего момента Mк(z)')
    add(d,'Схема и эпюра выполнены в масштабе по длине. Скачки показаны в местах приложения сосредоточенных крутящих моментов, наклоны — при действии распределённых моментов.')
    p=d.add_paragraph();p.alignment=center
    p.add_run().add_picture(io.BytesIO(teacher_png_bytes(c)),width=Cm(16.7))
    pg(d)
    h(d,'8. Подбор сечения из условия прочности')
    add(d,'Условие прочности: τmax = |Mк|max/Wρ ≤ [τ] = '+fmt(s.tau_allow_MPa)+' МПа.')
    add(d,'Для круглого сечения Wρ ≈ 0,2·d³, где d — диаметр в миллиметрах.')
    add(d,'dтр = ³√(|Mк|max·10⁶/(0,2·[τ])) = '+fmt(des['required_solid'])+' мм.')
    add(d,'По нормальному ряду принимаем dкр = '+fmt(des['d_round'],0)+' мм.')
    add(d,'τкр = '+fmt(c.max_T_knm)+'·10⁶/(0,2·'+fmt(des['d_round'],0)+'³) = '+fmt(des['tau_solid'])+' МПа ≤ '+fmt(s.tau_allow_MPa)+' МПа.')
    add(d,'Для кольцевого сечения Wρ ≈ 0,2·D³·(1 − (d/D)⁴), d/D = 0,5.')
    add(d,'Dтр = ³√(|Mк|max·10⁶/(0,2·[τ]·(1 − 0,5⁴))) = '+fmt(des['required_hollow'])+' мм.')
    add(d,'Принимаем D = '+fmt(des['D_hollow'],0)+' мм, d = D/2 = '+fmt(des['d_hollow'],0)+' мм (наружный размер округлён вверх, внутренний получается целым числом миллиметров).')
    add(d,'τкол = '+fmt(des['tau_hollow'])+' МПа ≤ '+fmt(s.tau_allow_MPa)+' МПа; условие прочности выполняется.')
    h(d,'9. Сравнение веса выбранных валов')
    add(d,'При одинаковой длине и плотности вес пропорционален площади поперечного сечения.')
    add(d,'Fкр = π·dкр²/4 = '+fmt(des['area_solid'])+' мм².')
    add(d,'Fкол = π·(D² − d²)/4 = '+fmt(des['area_hollow'])+' мм².')
    change=('снижение' if des['mass_ratio']<1 else 'увеличение')
    add(d,'Gкол/Gкр = Fкол/Fкр = '+fmt(des['mass_ratio'],4)+'. Относительное '+change+' веса кольцевого сечения: '+fmt(abs(des['weight_saved_pct']))+' %.')
    h(d,'Ответ:')
    comparison=('кольцевое сечение легче' if des['mass_ratio']<1 else ('круглое сечение легче' if des['mass_ratio']>1 else 'веса одинаковы'))
    add(d,'|Mк|max = '+fmt(c.max_T_knm)+' кН·м; круглый вал dкр = '+fmt(des['d_round'],0)+' мм; кольцевой вал D = '+fmt(des['D_hollow'],0)+' мм, d = '+fmt(des['d_hollow'],0)+' мм. Оба условия прочности выполняются; '+comparison+' при одинаковой длине и материале.')
    add(d,'Дополнительно: при введённом G = '+fmt(s.G_MPa)+' МПа программа рассчитывает углы закручивания и проверку жёсткости, но это не обязательная часть задания №2 по предоставленному образцу.',space=0)
    d.save(path)
    return path
