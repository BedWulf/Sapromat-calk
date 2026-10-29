"""Регрессия Word: напечатанные формулы согласованы с расчётом."""
import os
import tempfile
import unittest
from docx import Document
from generate_regressions import generate

class TestsWordExport(unittest.TestCase):
    def test_teacher_formula_compensation_is_on_third_segment(self):
        with tempfile.TemporaryDirectory() as d:
            path=os.path.join(d,'teacher.docx')
            result=generate(path,22,15,15)
            texts=[p.text.strip() for p in Document(path).paragraphs]
            eqs=[t for t in texts if t.startswith('N(z) =')]
            self.assertEqual(len(eqs),6)
            self.assertIn('‹z − 2l›¹',eqs[0])  # Универсальная формула с Маколеем.
            self.assertIn('(z − 2l)III',eqs[1])  # Вклад начинается в III.
            self.assertNotIn('(z − 2l)',eqs[3]) # На участке II компенсации НЕТ.
            self.assertIn('(z − 2l)',eqs[4])
            self.assertEqual(result['b'],12)
            self.assertEqual(result['h'],24)

    def test_nika_standard_size_and_stress_are_in_docx(self):
        with tempfile.TemporaryDirectory() as d:
            path=os.path.join(d,'nika.docx')
            result=generate(path,17,10,24)
            texts='\n'.join(p.text for p in Document(path).paragraphs)
            self.assertIn('b_min =',texts)
            self.assertIn('10.308 мм',texts)
            self.assertIn('b = 11 мм, h = 22 мм',texts)
            self.assertIn('140.50 МПа',texts)
            self.assertEqual(result['b'],11)

if __name__=='__main__':
    unittest.main()
