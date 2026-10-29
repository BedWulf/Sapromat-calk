import os,sys,tempfile,unittest,math
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT));sys.path.insert(0,str(ROOT/'torsion'))
from teacher_exact import create_tension_docx,create_torsion_docx
from model import Calculation
from teacher import teacher_settings,example_teacher,DATA,SCHEMES

class TestCombined(unittest.TestCase):
    def test_twenty_templates(self):
        from docx import Document
        c=Calculation(example_teacher())
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'torsion.docx'
            v=create_torsion_docx(path,c,{'variant':'11-2-23-4'})
            self.assertEqual(v['eight'],[5,5,-5,10,10,10,10,0]);self.assertEqual(v['diameters']['d_round'],85)
            d=Document(path);orig=Document(ROOT/'teacher_templates'/'teacher_torsion.docx')
            self.assertEqual(len(d.tables),len(orig.tables));self.assertEqual(len(d.paragraphs),len(orig.paragraphs))
            self.assertEqual(d.tables[1].cell(2,3).text,'-5')
            self.assertTrue(path.stat().st_size>20000)
            # Native OMML equations are kept (no false 180 MPa).
            ns='{http://schemas.openxmlformats.org/officeDocument/2006/math}t'
            texts=[n.text for n in d.paragraphs[75]._p.iter() if n.tag==ns]
            self.assertIn('=90 МПа',texts)
    def test_900_standard_torsion(self):
        for row in DATA:
            for no in SCHEMES:
                c=Calculation(teacher_settings(row,no))
                self.assertAlmostEqual(c.reaction_knm+c.applied_sum,0,places=6)
                self.assertTrue(math.isfinite(c.max_T_knm))
    def test_tension_original(self):
        sys.path.insert(0,str(ROOT/'tension'))
        from core.variants import get_preset_scheme
        fixed,loads=get_preset_scheme(22,1)
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'tension.docx'
            values=create_tension_docx(path,loads,{'P':15,'q':15,'l':1,'sigma_allow':160e6,'E':200e9},{'variant':'11-2-23-4'},fixed)
            self.assertEqual(values['eight'],[45,45,45,15,15,15,30,30])
            self.assertEqual((values['b'],values['h']),(12,24));self.assertAlmostEqual(values['sigma'],156.25)
            from docx import Document
            d=Document(path);orig=Document(ROOT/'teacher_templates'/'teacher_tension.docx')
            self.assertEqual(len(d.paragraphs),len(orig.paragraphs))
            self.assertEqual(len(d.tables),len(orig.tables))
            ns='{http://schemas.openxmlformats.org/officeDocument/2006/math}t'
            self.assertIn('=160 МПа',[n.text for n in d.paragraphs[67]._p.iter() if n.tag==ns])

if __name__=='__main__':unittest.main()
